#!/usr/bin/env python3
"""
build_xlsx.py — regenerate BPH_Master.xlsx (single-sheet mirror) from BPH_Master.csv.

The CSV is the source of truth; the workbook is a convenience copy for analysts who
live in Excel. Run this after every CSV change so the two never drift (the Sept-2026
audit found the workbook one row behind the CSV).

Formatting matches the established workbook: one sheet named "BPH_Master", header row
in bold white on dark slate (#1F2937), top-aligned and wrapped, frozen below the
header, fixed column widths; every data cell is written as a plain string (empty CSV
fields stay empty).

Usage:
  python tools/build_xlsx.py [--repo-root DIR]
Requires: openpyxl (pip install openpyxl)
Exit codes: 0 ok · 4 input/setup error
"""
from __future__ import annotations

import argparse
import csv
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_ROOT = os.path.dirname(HERE)

HEADER_FILL = "1F2937"
HEADER_FONT = "FFFFFF"
DEFAULT_WIDTH = 18
# Widths keyed by column name so a column reorder cannot misapply them.
WIDTHS = {
    "provider_name": 30,
    "status": 11,
    "risk_tier": 24,
    "primary_asn": 16,
    "additional_asns": 20,
    "sanctions_designations": 34,
    "associated_malware": 34,
    "notes": 60,
    "description": 50,
}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo-root", default=DEFAULT_ROOT)
    args = ap.parse_args()
    root = os.path.abspath(args.repo_root)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.stderr.write("[ABORT] openpyxl is required: pip install openpyxl\n")
        return 4

    csv_path = os.path.join(root, "BPH_Master.csv")
    if not os.path.exists(csv_path):
        sys.stderr.write(f"[ABORT] BPH_Master.csv not found: {csv_path}\n")
        return 4
    with open(csv_path, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        sys.stderr.write("[ABORT] BPH_Master.csv is empty\n")
        return 4
    header, width = rows[0], len(rows[0])
    for n, r in enumerate(rows[1:], start=2):
        if len(r) != width:
            sys.stderr.write(f"[ABORT] CSV line {n} has {len(r)} fields, header has {width}\n")
            return 4

    wb = Workbook()
    ws = wb.active
    ws.title = "BPH_Master"
    for r in rows:
        ws.append([v if v != "" else None for v in r])

    fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    font = Font(bold=True, color=HEADER_FONT)
    align = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = fill, font, align
    ws.freeze_panes = "A2"
    for i, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(name, DEFAULT_WIDTH)

    out = os.path.join(root, "BPH_Master.xlsx")
    wb.save(out)
    print(f"[ok] wrote {out} ({len(rows) - 1} data rows, {width} columns)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
